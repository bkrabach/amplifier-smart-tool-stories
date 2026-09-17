"""Deterministic editable PowerPoint adaptation of retained HTML slides."""

import io
import re
from collections import Counter

from .artifacts import parse_html, preview
from .errors import require

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
LIMITS = [
    "PowerPoint uses a semantic layout, not an exact reproduction of HTML styling.",
    "Text and tables are editable; images, SVG, charts, animation and arbitrary CSS are unsupported.",
    "Structural and text-coverage checks do not establish visual quality in PowerPoint.",
    "The dashboard previews the HTML source revision, not the converted PowerPoint.",
]


def words(text):
    return Counter(re.findall(r"\w+", text.casefold()))


def slide_text(slide):
    return " ".join(
        shape.text
        if shape.has_text_frame
        else " ".join(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.has_table
        else ""
        for shape in slide.shapes
    )


def convert(revision, title):
    from pptx import Presentation

    from .converters.powerpoint import convert_slides

    original = parse_html(revision["html"])
    require(
        not original.select("img,svg,canvas,video,audio,math,iframe,object,embed"),
        "PowerPoint conversion supports text and tables; remove unsupported media or use HTML export.",
        "unsupported_content",
    )
    clean, _ = preview(revision["html"])
    soup = parse_html(clean)
    source_slides = soup.select(".slide") or soup.body.find_all("section", recursive=False)
    require(0 < len(source_slides) <= 12, "PowerPoint requires 1–12 slide sections.", "unsupported_content")
    # Content outside slide sections would otherwise silently disappear.
    body = parse_html(clean).body
    for node in body.select(".slide") or body.select("section"):
        node.decompose()
    for node in body.select("style,nav,.navigation,.slide-counter,.nav-dots"):
        node.decompose()
    require(
        not body.get_text(strip=True),
        "Move all presentation text inside slide sections.",
        "unsupported_content",
    )
    prs = convert_slides(source_slides)
    require(len(prs.slides) == len(source_slides), "Conversion lost slides.", "conversion_failed")
    for index, (source, slide) in enumerate(zip(source_slides, prs.slides), 1):
        for node in source.select("style,script"):
            node.decompose()
        missing = words(source.get_text(" ")) - words(slide_text(slide))
        require(
            not missing,
            f"Slide {index}: conversion omitted text: {', '.join(missing)[:200]}",
            "conversion_failed",
        )
        for shape in slide.shapes:
            require(
                shape.left >= 0
                and shape.top >= 0
                and shape.left + shape.width <= prs.slide_width + 9144
                and shape.top + shape.height <= prs.slide_height + 9144,
                f"Slide {index}: converted content exceeds the slide canvas.",
                "conversion_failed",
            )
        notes = [f"Source revision: {revision['id']}", f"HTML SHA-256: {revision['sha256']}"]
        for fact in revision["evidence"]:
            notes.append(f"{fact['id']} — {fact['source_id']}: {fact['quote']}")
        notes.extend(revision.get("limitations", []))
        slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
    prs.core_properties.title = title
    prs.core_properties.subject = f"Derived from Stories revision {revision['id']}"
    buffer = io.BytesIO()
    prs.save(buffer)
    data = buffer.getvalue()
    reopened = Presentation(io.BytesIO(data))
    require(len(reopened.slides) == len(source_slides), "PowerPoint round-trip failed.", "conversion_failed")
    return data, {
        "slide_count": len(prs.slides),
        "checks": {
            "structural": "passed",
            "text_coverage": "passed",
            "shape_bounds": "passed",
            "visual": "not_performed",
        },
        "limitations": LIMITS,
    }
