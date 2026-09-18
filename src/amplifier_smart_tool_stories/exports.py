"""Explicit document exports. No source paths, model calls or review overlays."""

import asyncio
import base64
import hashlib
import io

from .errors import require


def word(revision):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    value = revision["document"]
    doc = Document()
    doc.core_properties.title = value["title"]
    doc.core_properties.subject = "Stories revision " + revision["id"]
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(2 / 3)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Georgia", Pt(12.75)
    normal.font.color.rgb = RGBColor.from_string("252C2A")
    normal.paragraph_format.line_spacing = 1.55
    normal.paragraph_format.space_after = Pt(13.5)
    normal.paragraph_format.keep_together = True
    for name, size in [("Title", 27), ("Heading 1", 17.25), ("Subtitle", 10.5)]:
        style = doc.styles[name]
        style.font.name = "Arial" if name == "Subtitle" else "Georgia"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("64716B" if name == "Subtitle" else "252C2A")
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.keep_with_next = name != "Subtitle"
        style.paragraph_format.space_before = Pt(21 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(9 if name == "Heading 1" else 13.5)
    # Word theme font aliases and title borders must not override the document design.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn("w:rFonts")):
            for key in list(fonts.attrib):
                if "theme" in key.lower():
                    del fonts.attrib[key]
    doc.styles["Title"].font.bold = False
    doc.styles["Title"].font.italic = False
    doc.styles["Subtitle"].font.italic = False
    for name in ("Title", "Subtitle", "Heading 1"):
        for spacing in list(doc.styles[name].element.iter(qn("w:spacing"))):
            if spacing.getparent().tag == qn("w:rPr"):
                spacing.getparent().remove(spacing)
    doc.add_paragraph(value["title"], "Title")
    if value["subtitle"]:
        doc.add_paragraph(value["subtitle"], "Subtitle")
    for b in value["blocks"]:
        citations = " [" + ", ".join(b["evidence_ids"]) + "]" if b["evidence_ids"] else ""
        if b["kind"] == "table":
            if b["text"] or citations:
                caption = doc.add_paragraph(b["text"] + citations)
                caption.paragraph_format.keep_with_next = True
            table = doc.add_table(rows=0, cols=len(b["rows"][0]))
            table.autofit = False
            for i, row in enumerate(b["rows"]):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = 1.45
                        paragraph.paragraph_format.keep_with_next = i < len(b["rows"]) - 1
                        paragraph.paragraph_format.space_after = Pt(7.5)
                        for run in paragraph.runs:
                            run.font.size = Pt(10.5)
                            run.font.bold = i == 0
                    props = cell._tc.get_or_add_tcPr()
                    if i == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "F1F5F1")
                        props.append(shade)
                rowprops = table.rows[-1]._tr.get_or_add_trPr()
                rowprops.append(OxmlElement("w:cantSplit"))
                if i == 0:
                    rowprops.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
        elif b["kind"] == "list":
            if b["text"]:
                doc.add_paragraph(b["text"])
            for i, item in enumerate(b["items"]):
                paragraph = doc.add_paragraph(
                    item + (citations if i == len(b["items"]) - 1 else ""), "List Bullet"
                )
                paragraph.paragraph_format.space_after = Pt(6)
        else:
            paragraph = doc.add_paragraph(b["text"], "Heading 1" if b["kind"] == "heading" else "Normal")
            if b["kind"] == "quote":
                paragraph.paragraph_format.left_indent = Pt(13.5)
                paragraph.paragraph_format.right_indent = Pt(13.5)
            if citations:
                run = paragraph.add_run(citations)
                run.font.name = "Arial"
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor.from_string("3C5649")
    if revision["evidence"]:
        used = {x for b in value["blocks"] for x in b["evidence_ids"]}
        doc.add_paragraph("Source references", "Heading 1")
        groups = {}
        for fact in revision["evidence"]:
            if fact["id"] in used:
                groups.setdefault(fact["source_id"], []).append(fact["id"])
        for source, ids in groups.items():
            doc.add_paragraph(f"{source} — [{', '.join(ids)}]")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def artifact(revision, format, media_content=None):
    require(
        isinstance(format, str) and format in {"html", "zip", "pdf", "docx"},
        "Choose html, zip, pdf or docx.",
        "unsupported_format",
    )
    limitations = []
    checks = {"structure": "passed", "visual": "not_performed"}
    if format in {"html", "zip"}:
        from .artifacts import parse_html
        from .media import EXTENSIONS, bindings

        assets = revision.get("assets", [])
        if format == "zip" or assets:
            from .media import portable_markup

            portable_markup(revision["html"])
        content = media_content or {}
        used, missing = bindings(revision["html"], assets, strict=format == "zip" or bool(assets))
        selected = [a for a in assets if a["id"] in used]
        for asset in selected:
            raw = content.get(asset["id"])
            require(
                raw is not None and hashlib.sha256(raw).hexdigest() == asset["sha256"],
                "Required media is missing or changed.",
                "missing_asset",
            )
        require(
            format != "html" or not any(not a["mime_type"].startswith("image/") for a in selected),
            "This revision includes video or captions. Choose ZIP to deliver HTML with its media assets.",
            "package_required",
        )
        markup = revision["html"]
        paths = {a["id"]: "assets/" + a["id"] + "." + EXTENSIONS[a["mime_type"]] for a in selected}
        if selected:
            soup = parse_html(markup)
            for node in soup.select("img,video,source,track"):
                for attr in ("src", "poster"):
                    ref = node.get(attr, "")
                    if ref.startswith("asset:") and ref[6:] in paths:
                        asset = next(a for a in selected if a["id"] == ref[6:])
                        node[attr] = (
                            paths[asset["id"]]
                            if format == "zip"
                            else (
                                "data:"
                                + asset["mime_type"]
                                + ";base64,"
                                + base64.b64encode(content[asset["id"]]).decode()
                            )
                        )
            markup = str(soup)
        # Review navigation belongs to the dashboard; portable decks need their own player.
        # Imported scripted decks retain their existing presentation behavior.
        soup = parse_html(markup)
        if revision.get("kind") != "document" and soup.select(".slide") and not soup.find("script"):
            from importlib.resources import files

            player = soup.new_tag("script")
            player["data-stories-presentation"] = "1"
            player.string = (
                files("amplifier_smart_tool_stories").joinpath("resources/presentation.js").read_text()
            )
            (soup.body or soup).append(player)
            markup = str(soup)
        data, mime = markup.encode(), "text/html"
        checks = {**revision["review"], "media_playback": "not_performed"}
        if missing:
            limitations.append(
                "Imported HTML has unretained media references; external files are not packaged."
            )
        if format == "zip":
            import json
            import zipfile

            output = io.BytesIO()
            with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as archive:

                def write_member(name, value):
                    archive.writestr(zipfile.ZipInfo(name), value)

                write_member("index.html", data)
                if revision.get("storyboard") is not None:
                    write_member(
                        "storyboard.json",
                        json.dumps(
                            {
                                "schema_version": 1,
                                "revision_id": revision["id"],
                                "direction_id": revision["direction_id"],
                                "brief_id": revision["brief_id"],
                                "brief": revision["brief"],
                                "audience": revision["audience"],
                                "fidelity": revision["fidelity"],
                                "storyboard": revision["storyboard"],
                                "evidence": revision["evidence"],
                                "assets": selected,
                                "asset_paths": paths,
                            },
                            ensure_ascii=False,
                            indent=2,
                        ),
                    )
                for asset in selected:
                    write_member(paths[asset["id"]], content[asset["id"]])
                write_member(
                    "manifest.json",
                    json.dumps(
                        {
                            "revision_id": revision["id"],
                            "delivery_sha256": revision.get("delivery_sha256", revision["sha256"]),
                            "assets": selected,
                        },
                        indent=2,
                    ),
                )
                write_member(
                    "README.txt",
                    "Extract this ZIP, then open index.html in a browser. Keep the assets folder beside it. No Stories service is needed.\nVideo codec support depends on the browser.\n",
                )
            data, mime = output.getvalue(), "application/zip"
        if selected:
            limitations.append(
                "Static layout review does not certify video playback or audio. Browser codec support varies."
            )
        if revision.get("kind") == "storyboard" and format == "html":
            limitations.append(
                "Review-only HTML. Choose ZIP to include editable storyboard.json and exact assets."
            )
    else:
        require(
            revision.get("document") is not None,
            "PDF/Word export supports structured documents only.",
            "unsupported_format",
        )
        if format == "docx":
            data, mime = (
                word(revision),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            limitations = [
                "Editable Word document; line and page breaks depend on the Word renderer and installed fonts.",
                "This export has not received per-artifact visual review. HTML review does not certify Word layout.",
            ]
        else:
            from .quality import render

            result = asyncio.run(render(revision["html"], 40, include_pdf=True))
            require(
                not result["findings"],
                "PDF layout has findings: " + "; ".join(result["findings"]),
                "export_review_failed",
            )
            data, mime = base64.b64decode(result["pdf"]), "application/pdf"
            checks = {
                "structure": "passed",
                "text_coverage": "passed",
                "bounds": "passed",
                "visual": "not_performed",
            }
            limitations = [
                "Letter PDF from the retained HTML. Browser pagination is approximate; inspect the PDF before delivery."
            ]
    return {
        "revision_id": revision["id"],
        "format": format,
        "mime_type": mime,
        "sha256": hashlib.sha256(data).hexdigest(),
        "source_sha256": revision["sha256"],
        "data_base64": base64.b64encode(data).decode(),
        "size_bytes": len(data),
        "delivery_sha256": revision.get("delivery_sha256", revision["sha256"]),
        "media_warnings": revision.get("media_warnings", []),
        "checks": checks,
        "limitations": limitations,
    }
