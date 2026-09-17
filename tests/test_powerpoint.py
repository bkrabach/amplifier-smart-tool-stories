import base64
import io
import json
import os
import subprocess
import sys

import pytest
from pptx import Presentation

from amplifier_smart_tool_stories import Stories, StoriesError

HTML = """<html><body><section class="slide"><h1>Benchmark results</h1>
<p>Median duration fell from 10 seconds to 8 seconds in one benchmark only.</p>
<table><tr><th>Run</th><th>Seconds</th></tr><tr><td>Before</td><td>10</td></tr>
<tr><td>After</td><td>8</td></tr></table></section>
<section class="slide"><h1>Limitations</h1><ul><li>No general productivity claim.</li>
<li>Sample size remains unknown.</li></ul></section></body></html>"""


def setup(tmp_path, html=HTML):
    api = Stories(tmp_path / "state")
    r = api.create_story("Qualified result", html, "import")
    return api, r["story_id"], r["revision_id"]


def test_editable_roundtrip_revision_and_overlay_exclusion(tmp_path):
    api, sid, rid = setup(tmp_path)
    api.add_comment(sid, rid, "PRIVATE REVIEW COMMENT", "comment", {"kind": "story"}, "agent")
    before = api.get_story(sid)
    result = api.export(sid, rid, str(tmp_path / "deck.pptx"), "pptx")
    prs = Presentation(result["path"])
    assert len(prs.slides) == 2
    assert any(s.has_table for s in prs.slides[0].shapes)
    assert all(not s.shape_type == 13 for sl in prs.slides for s in sl.shapes)
    assert any("one benchmark only" in " ".join(s.text.split()) for s in prs.slides[0].shapes if s.has_text_frame)
    assert rid in prs.slides[0].notes_slide.notes_text_frame.text
    assert "PRIVATE REVIEW COMMENT" not in prs.slides[0].notes_slide.notes_text_frame.text
    assert result["source_sha256"] == api.get_revision(sid, rid)["sha256"]
    assert result["checks"]["visual"] == "not_performed"
    assert api.get_story(sid) == before
    with pytest.raises(StoriesError, match="already exists"):
        api.export(sid, rid, result["path"], "pptx")
    assert base64.b64decode(api.get_export(sid, rid)["data"]).decode() == HTML


@pytest.mark.parametrize(
    "body",
    [
        '<section class="slide"><h1>Image</h1><img src="file:///private/image.png"></section>',
        "<p>Not a slide</p>",
        '<p>Outside slide</p><section class="slide"><h1>Inside</h1></section>',
        '<section class="slide"><h1>Too much</h1><p>' + "Words retained. " * 800 + "</p></section>",
        '<section class="slide"><h1>Merged</h1><table><tr><td colspan="2">X</td></tr></table></section>',
    ],
)
def test_unsupported_or_overfull_artifacts_never_write_output(tmp_path, body):
    api, sid, rid = setup(tmp_path, "<html><body>" + body + "</body></html>")
    output = tmp_path / "bad.pptx"
    with pytest.raises(StoriesError):
        api.export(sid, rid, str(output), "pptx")
    assert not output.exists()


def test_cli_export_without_credentials_from_unrelated_directory(tmp_path):
    api, sid, rid = setup(tmp_path)
    env = {k: v for k, v in os.environ.items() if not any(s in k for s in ("KEY", "TOKEN", "SECRET"))}
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "amplifier_smart_tool_stories",
            "--store",
            str(tmp_path / "state"),
            "export",
            "--input",
            json.dumps(
                {
                    "story_id": sid,
                    "revision_id": rid,
                    "output_path": str(tmp_path / "cli.pptx"),
                    "format": "pptx",
                }
            ),
        ],
        cwd=tmp_path,
        env=env,
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert len(Presentation(json.loads(result.stdout)["path"]).slides) == 2


def test_evidence_quotes_survive_in_notes(tmp_path):
    from amplifier_smart_tool_stories.powerpoint import convert

    revision = {
        "id": "r1",
        "html": HTML,
        "sha256": "source-hash",
        "limitations": ["One benchmark only"],
        "evidence": [{"id": "f1", "source_id": "s1", "quote": "10 seconds to 8 seconds"}],
    }
    data, _ = convert(revision, "Title")
    prs = Presentation(io.BytesIO(data))
    assert "f1 — s1: 10 seconds to 8 seconds" in prs.slides[0].notes_slide.notes_text_frame.text
